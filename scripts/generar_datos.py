"""Genera los datos de ejemplo de Atiende Primero. TODO ES FICTICIO.

Produce, con una semilla fija (mismo resultado en cada ejecución):
  datos/reportes.xlsx        hojas "Reportes" y "Acerca de"
  datos/criterios.xlsx       hojas "Pesos", "Umbrales", "Reglas" y "Acerca de"
  sitio/descargas/*.xlsx     copias para el botón "Descargar Excel de ejemplo"
  sitio/datos/ejemplo.js     el contenido de ambos Excel, para que el sitio abra con datos sin cargar nada

Los reportes, folios, cuadrillas y cifras son inventados. Solo se usan datos públicos del INEGI para
repartir los reportes según la población y ubicarlos cerca de localidades reales (ver docs/mapa.md).

Uso: python3 scripts/generar_datos.py
"""

import csv
import json
import math
import random
import shutil
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------------------
# Parámetros generales
# ---------------------------------------------------------------------------
SEMILLA = 20261020
TOTAL_REPORTES = 200
FECHA_CORTE = datetime(2026, 10, 19, 8, 0)   # las horas abiertas se miden contra esta fecha
EXPONENTE_POBLACION = 0.75                    # < 1 reparte un poco más hacia municipios pequeños
NUM_CUADRILLAS = 20
AVISO = "Ejemplo didáctico con datos ficticios. No corresponde a ninguna institución ni a reportes reales."

RAIZ = Path(__file__).resolve().parent.parent
REFERENCIA = RAIZ / "scripts" / "referencia"
MAPA_JS = RAIZ / "sitio" / "datos" / "municipios.js"

# Tipos de falla: plazo de atención en horas, severidades posibles, usuarios afectados (mínimo, típico, máximo)
# y probabilidad de que el reporte indique riesgo a la seguridad. Valores inventados, no son normas oficiales.
TIPOS = {
    "Sin servicio":          {"peso": 30, "plazo": 8,  "sev": [2, 3, 3, 4],    "usuarios": (5, 120, 2500), "riesgo": 0.03},
    "Variación de voltaje":  {"peso": 18, "plazo": 24, "sev": [1, 2, 2, 3],    "usuarios": (2, 40, 600),   "riesgo": 0.05},
    "Poste dañado":          {"peso": 12, "plazo": 48, "sev": [2, 2, 3, 3, 4], "usuarios": (0, 25, 300),   "riesgo": 0.20},
    "Cable caído":           {"peso": 5,  "plazo": 4,  "sev": [3, 3, 4, 4],    "usuarios": (1, 30, 400),   "riesgo": 0.85},
    "Transformador dañado":  {"peso": 10, "plazo": 24, "sev": [3, 3, 4],       "usuarios": (15, 80, 400),  "riesgo": 0.15},
    "Alumbrado":             {"peso": 14, "plazo": 72, "sev": [1, 1, 2],       "usuarios": (0, 0, 0),      "riesgo": 0.04},
    "Otro":                  {"peso": 8,  "plazo": 72, "sev": [1, 2, 2, 3],    "usuarios": (0, 5, 80),     "riesgo": 0.03},
}
ESTADOS = [("Abierto", 0.55), ("En proceso", 0.30), ("Cerrado", 0.15)]

COLUMNAS = [
    "Folio", "Fecha y hora de alta", "Municipio", "Localidad", "Latitud", "Longitud",
    "Tipo de falla", "Severidad", "Usuarios afectados", "Horas abiertas", "Plazo de atención (horas)",
    "Estado", "Riesgo a la seguridad", "Reportes previos (30 días)", "Cuadrilla",
]

# ---------------------------------------------------------------------------
# Criterios iniciales (la Fase 2 define la fórmula; aquí solo se dejan los valores editables)
# ---------------------------------------------------------------------------
PESOS = [
    ["Criterio", "Descripción", "Peso", "Eje de la matriz"],
    ["Severidad", "Gravedad técnica de la falla, de 1 (menor) a 4 (mayor).", 30, "Impacto"],
    ["Usuarios afectados", "Cuántos usuarios se quedan sin servicio o con servicio deficiente (estimado).", 25, "Impacto"],
    ["Tiempo respecto al plazo", "Horas abiertas comparadas con el plazo de atención; crece al acercarse o pasar el plazo.", 20, "Urgencia"],
    ["Recurrencia", "Reportes previos en la misma localidad en los últimos 30 días.", 10, "Urgencia"],
    ["Riesgo a la seguridad", "Si el reporte indica riesgo para personas o bienes.", 15, "Ninguno"],
]
UMBRALES = [
    ["Prioridad", "Puntaje mínimo", "Qué significa"],
    ["Crítica", 70, "Atender de inmediato."],
    ["Alta", 50, "Atender en el turno."],
    ["Media", 30, "Programar en los próximos días."],
    ["Baja", 0, "Atender cuando haya capacidad."],
]
REGLAS = [
    ["Regla", "Campo", "Condición", "Valor", "Prioridad mínima", "Activa"],
    ["Si el tipo de falla es cable caído, la prioridad es crítica.", "Tipo de falla", "es igual a", "Cable caído", "Crítica", "Sí"],
    ["Si hay riesgo a la seguridad, la prioridad es crítica.", "Riesgo a la seguridad", "es igual a", "Sí", "Crítica", "Sí"],
    ["Si la severidad es 4, la prioridad es al menos alta.", "Severidad", "mayor o igual que", 4, "Alta", "Sí"],
    ["Si hay 500 usuarios afectados o más, la prioridad es al menos alta.", "Usuarios afectados", "mayor o igual que", 500, "Alta", "Sí"],
]


# ---------------------------------------------------------------------------
# Geografía
# ---------------------------------------------------------------------------
def cargar_municipios():
    """Lee la geometría del sitio para usar los mismos nombres y comprobar que cada punto cae en su municipio."""
    texto = MAPA_JS.read_text(encoding="utf-8")
    geo = json.loads(texto[texto.index("{"): texto.rindex("}") + 1])
    municipios = {}
    for f in geo["features"]:
        g = f["geometry"]
        poligonos = g["coordinates"] if g["type"] == "MultiPolygon" else [g["coordinates"]]
        municipios[f["properties"]["clave"]] = {"nombre": f["properties"]["nombre"], "poligonos": poligonos}
    return municipios


def punto_en_anillo(lon, lat, anillo):
    dentro = False
    j = len(anillo) - 1
    for i in range(len(anillo)):
        xi, yi = anillo[i]
        xj, yj = anillo[j]
        if (yi > lat) != (yj > lat) and lon < (xj - xi) * (lat - yi) / (yj - yi) + xi:
            dentro = not dentro
        j = i
    return dentro


def punto_en_municipio(lon, lat, municipio):
    for poligono in municipio["poligonos"]:
        if punto_en_anillo(lon, lat, poligono[0]) and not any(punto_en_anillo(lon, lat, h) for h in poligono[1:]):
            return True
    return False


def cargar_referencias(municipios):
    poblacion = {r["clave"]: int(r["poblacion_2020"]) for r in csv.DictReader(open(REFERENCIA / "municipios_sonora.csv", encoding="utf-8"))}
    localidades = {}
    for r in csv.DictReader(open(REFERENCIA / "localidades_sonora.csv", encoding="utf-8")):
        lat, lon = float(r["latitud"]), float(r["longitud"])
        # Solo localidades cuyo punto cae dentro de su municipio en el mapa simplificado
        if punto_en_municipio(lon, lat, municipios[r["clave_mun"]]):
            localidades.setdefault(r["clave_mun"], []).append(
                {"nombre": r["localidad"], "lat": lat, "lon": lon, "poblacion": int(r["poblacion_2020"])})
    return poblacion, localidades


def ubicar(rng, localidad, municipio):
    """Coordenada aleatoria cerca de la localidad; el radio crece con el tamaño de la localidad."""
    radio_km = 0.4 + 0.0045 * math.sqrt(localidad["poblacion"])
    for _ in range(50):
        d = min(abs(rng.gauss(0, radio_km / 2)), radio_km)
        ang = rng.uniform(0, 2 * math.pi)
        lat = localidad["lat"] + (d / 111.0) * math.sin(ang)
        lon = localidad["lon"] + (d / (111.0 * math.cos(math.radians(localidad["lat"])))) * math.cos(ang)
        if punto_en_municipio(lon, lat, municipio):
            return round(lat, 5), round(lon, 5)
    return round(localidad["lat"], 5), round(localidad["lon"], 5)


# ---------------------------------------------------------------------------
# Reportes
# ---------------------------------------------------------------------------
def elegir(rng, opciones_pesos):
    opciones, pesos = zip(*opciones_pesos)
    return rng.choices(opciones, weights=pesos)[0]


def usuarios_afectados(rng, minimo, tipico, maximo, poblacion_localidad):
    if maximo == 0:
        return 0
    valor = int(round(rng.lognormvariate(math.log(max(tipico, 1)), 0.9)))
    tope = min(maximo, max(minimo, poblacion_localidad // 3))
    return max(minimo, min(valor, tope))


def horas_abiertas(rng, plazo):
    """Horas abiertas según la proporción respecto al plazo: ~70 % en plazo, ~20 % vencido, ~10 % muy vencido."""
    r = rng.random()
    if r < 0.70:
        proporcion = rng.uniform(0.05, 0.95)
    elif r < 0.90:
        proporcion = rng.uniform(1.0, 2.0)
    else:
        proporcion = rng.uniform(2.0, 4.0)
    return max(1, round(plazo * proporcion))


def reporte(rng, tipo, municipio, localidad, coords, **fijos):
    t = TIPOS[tipo]
    estado = fijos.get("estado") or elegir(rng, ESTADOS)
    horas = fijos.get("horas") or horas_abiertas(rng, t["plazo"])
    alta = FECHA_CORTE - timedelta(hours=horas, minutes=rng.randint(0, 59))
    if estado == "Cerrado":  # para los cerrados, "horas abiertas" es lo que tardó en cerrarse
        horas = max(1, round(horas * rng.uniform(0.3, 0.9)))
    if estado == "Abierto" and rng.random() < 0.5:
        cuadrilla = "Sin asignar"
    else:
        cuadrilla = f"Cuadrilla {rng.randint(1, NUM_CUADRILLAS):02d}"
    return {
        "Fecha y hora de alta": alta,
        "Municipio": municipio["nombre"],
        "Localidad": localidad["nombre"],
        "Latitud": coords[0],
        "Longitud": coords[1],
        "Tipo de falla": tipo,
        "Severidad": fijos.get("severidad", rng.choice(t["sev"])),
        "Usuarios afectados": fijos.get("usuarios", usuarios_afectados(rng, *t["usuarios"], localidad["poblacion"])),
        "Horas abiertas": horas,
        "Plazo de atención (horas)": t["plazo"],
        "Estado": estado,
        "Riesgo a la seguridad": fijos.get("riesgo", "Sí" if rng.random() < t["riesgo"] else "No"),
        "Reportes previos (30 días)": fijos.get("previos", min(rng.choices([0, 1, 2, 3], weights=[55, 25, 13, 7])[0], 3)),
        "Cuadrilla": fijos.get("cuadrilla", cuadrilla),
        "_caso": fijos.get("caso", ""),
    }


def buscar_localidad(localidades, clave_mun, nombre):
    return next(l for l in localidades[clave_mun] if l["nombre"] == nombre)


def casos_de_prueba(rng, municipios, localidades):
    """Casos incluidos a propósito para poner a prueba las reglas (ver docs/datos.md)."""
    casos = []

    def agregar(clave_mun, nombre_loc, tipo, **fijos):
        mun, loc = municipios[clave_mun], buscar_localidad(localidades, clave_mun, nombre_loc)
        casos.append(reporte(rng, tipo, mun, loc, ubicar(rng, loc, mun), **fijos))

    # 1. Cable caído con pocos usuarios: puntaje bajo, pero la regla dura lo vuelve crítico.
    agregar("26030", "Bahía de Kino", "Cable caído", severidad=2, usuarios=3, horas=1, previos=0,
            riesgo="No", estado="Abierto", cuadrilla="Sin asignar", caso="Cable caído con pocos usuarios")
    # 2. Riesgo a la seguridad en una falla menor: la regla de riesgo lo vuelve crítico.
    agregar("26043", "Heroica Nogales", "Variación de voltaje", severidad=1, usuarios=4, horas=2, previos=0,
            riesgo="Sí", estado="Abierto", caso="Riesgo a la seguridad en falla menor")
    # 3. Reportes vencidos respecto a su plazo.
    agregar("26055", "San Luis Río Colorado", "Alumbrado", severidad=1, usuarios=0, horas=230, previos=1,
            riesgo="No", estado="Abierto", caso="Vencido (más de 3 veces el plazo)")
    agregar("26042", "Navojoa", "Poste dañado", severidad=3, usuarios=40, horas=150, previos=0,
            riesgo="No", estado="En proceso", caso="Vencido (más de 3 veces el plazo)")
    agregar("26029", "Heroica Guaymas", "Sin servicio", severidad=3, usuarios=180, horas=20, previos=1,
            riesgo="No", estado="Abierto", caso="Vencido (más del doble del plazo)")
    # 4. Localidad con muchos reportes recurrentes: Pueblo Yaqui (Cajeme).
    for tipo, sev, usu, horas, previos in [
        ("Sin servicio", 3, 90, 3, 6), ("Variación de voltaje", 2, 25, 10, 7), ("Sin servicio", 2, 45, 5, 8),
        ("Transformador dañado", 3, 60, 12, 8), ("Variación de voltaje", 2, 15, 6, 9), ("Sin servicio", 3, 120, 2, 9),
    ]:
        agregar("26018", "Pueblo Yaqui", tipo, severidad=sev, usuarios=usu, horas=horas, previos=previos,
                riesgo="No", estado="Abierto", caso="Localidad recurrente")
    # 5. Falla muy grave dentro de plazo: alto impacto, baja urgencia.
    agregar("26030", "Hermosillo", "Sin servicio", severidad=4, usuarios=1800, horas=1, previos=0,
            riesgo="No", estado="Abierto", caso="Alto impacto, baja urgencia")
    # 6. Cable caído ya cerrado: no debe contar entre los pendientes.
    agregar("26018", "Ciudad Obregón", "Cable caído", severidad=4, usuarios=35, horas=3, previos=0,
            riesgo="Sí", estado="Cerrado", caso="Crítico pero cerrado")
    return casos


def generar_reportes(rng, municipios, poblacion, localidades):
    claves = [c for c in municipios if c in localidades]
    pesos_mun = [poblacion[c] ** EXPONENTE_POBLACION for c in claves]
    reportes = casos_de_prueba(rng, municipios, localidades)
    tipos = [(t, d["peso"]) for t, d in TIPOS.items()]
    while len(reportes) < TOTAL_REPORTES:
        clave = rng.choices(claves, weights=pesos_mun)[0]
        loc = rng.choices(localidades[clave], weights=[l["poblacion"] for l in localidades[clave]])[0]
        mun = municipios[clave]
        reportes.append(reporte(rng, elegir(rng, tipos), mun, loc, ubicar(rng, loc, mun)))
    reportes.sort(key=lambda r: r["Fecha y hora de alta"])
    for i, r in enumerate(reportes, start=1):
        r["Folio"] = f"REP-{i:04d}"
    return reportes


# ---------------------------------------------------------------------------
# Escritura de archivos
# ---------------------------------------------------------------------------
ENCABEZADO = Font(bold=True, color="FFFFFF")
FONDO = PatternFill("solid", fgColor="3B4A5A")
FIJA = datetime(2026, 10, 20, 0, 0)  # metadatos fijos para que el archivo no cambie en cada ejecución


def hoja_tabla(ws, filas, anchos=None):
    for fila in filas:
        ws.append(fila)
    for celda in ws[1]:
        celda.font, celda.fill = ENCABEZADO, FONDO
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for i, ancho in enumerate(anchos or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def filas_acerca(extra):
    return [["Dato", "Valor"], ["Aviso", AVISO]] + extra + [
        ["Generado con", "scripts/generar_datos.py (semilla fija %d)" % SEMILLA],
        ["Documentación", "docs/datos.md"],
    ]


ACERCA_REPORTES = filas_acerca([
    ["Fecha de corte", FECHA_CORTE],
    ["Qué es la fecha de corte", "Momento contra el que se miden las horas abiertas. Para los reportes cerrados, "
                                 "las horas abiertas son lo que tardó en cerrarse."],
])
ACERCA_CRITERIOS = filas_acerca([
    ["Pesos", "Cuánto aporta cada criterio al puntaje. Se recomienda que sumen 100."],
    ["Umbrales", "Puntaje mínimo para cada prioridad."],
    ["Reglas", "Reglas duras: el puntaje nunca puede dejar un reporte por debajo de la prioridad mínima indicada."],
])


def lista_desplegable(ws, columna, opciones, filas):
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(opciones), allow_blank=False)
    dv.error, dv.errorTitle = "Elige un valor de la lista.", "Valor no válido"
    ws.add_data_validation(dv)
    dv.add(f"{columna}2:{columna}{filas + 1}")


def fijar_metadatos(wb):
    wb.properties.creator = "Atiende Primero (datos ficticios)"
    wb.properties.created = FIJA
    wb.properties.modified = FIJA


def escribir_reportes(reportes, ruta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Reportes"
    hoja_tabla(ws, [COLUMNAS] + [[r[c] for c in COLUMNAS] for r in reportes],
               [11, 18, 22, 26, 10, 11, 21, 10, 11, 10, 12, 11, 12, 13, 13])
    for celda in ws["B"][1:]:
        celda.number_format = "yyyy-mm-dd hh:mm"
    n = len(reportes)
    lista_desplegable(ws, "G", list(TIPOS), n)
    lista_desplegable(ws, "L", [e for e, _ in ESTADOS], n)
    lista_desplegable(ws, "M", ["Sí", "No"], n)
    hoja_tabla(wb.create_sheet("Acerca de"), ACERCA_REPORTES, [28, 90])
    wb["Acerca de"]["B3"].number_format = "yyyy-mm-dd hh:mm"
    fijar_metadatos(wb)
    wb.save(ruta)


def escribir_criterios(ruta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Pesos"
    hoja_tabla(ws, PESOS, [26, 80, 8, 16])
    hoja_tabla(wb.create_sheet("Umbrales"), UMBRALES, [12, 16, 40])
    ws_reglas = wb.create_sheet("Reglas")
    hoja_tabla(ws_reglas, REGLAS, [62, 24, 20, 14, 18, 8])
    lista_desplegable(ws_reglas, "C", ["es igual a", "mayor o igual que", "menor o igual que"], 50)
    lista_desplegable(ws_reglas, "E", ["Crítica", "Alta", "Media", "Baja"], 50)
    lista_desplegable(ws_reglas, "F", ["Sí", "No"], 50)
    hoja_tabla(wb.create_sheet("Acerca de"), ACERCA_CRITERIOS, [28, 90])
    fijar_metadatos(wb)
    wb.save(ruta)


def a_texto(valor):
    return valor.strftime("%Y-%m-%d %H:%M") if isinstance(valor, datetime) else valor


def escribir_js(reportes, ruta):
    """Guarda el contenido de ambos Excel tal cual (hoja -> filas) para que el sitio lo procese igual que un archivo subido."""
    datos = {
        "aviso": AVISO,
        "reportes": {
            "Reportes": [COLUMNAS] + [[a_texto(r[c]) for c in COLUMNAS] for r in reportes],
            "Acerca de": [[a_texto(v) for v in fila] for fila in ACERCA_REPORTES],
        },
        "criterios": {"Pesos": PESOS, "Umbrales": UMBRALES, "Reglas": REGLAS, "Acerca de": ACERCA_CRITERIOS},
    }
    with open(ruta, "w", encoding="utf-8") as s:
        s.write("// Datos de ejemplo de Atiende Primero. TODO ES FICTICIO.\n")
        s.write("// Generado con scripts/generar_datos.py. No editar a mano.\n")
        s.write("window.DATOS_EJEMPLO = ")
        json.dump(datos, s, ensure_ascii=False, separators=(",", ":"))
        s.write(";\n")


def resumen(reportes):
    from collections import Counter
    print(f"Reportes: {len(reportes)}  ·  fecha de corte {FECHA_CORTE:%Y-%m-%d %H:%M}")
    print("Estados:", dict(Counter(r["Estado"] for r in reportes)))
    print("Tipos:", dict(Counter(r["Tipo de falla"] for r in reportes)))
    print("Municipios con reportes:", len({r["Municipio"] for r in reportes}))
    print("Top municipios:", Counter(r["Municipio"] for r in reportes).most_common(8))
    abiertos = [r for r in reportes if r["Estado"] != "Cerrado"]
    vencidos = [r for r in abiertos if r["Horas abiertas"] > r["Plazo de atención (horas)"]]
    print(f"Pendientes: {len(abiertos)}  ·  vencidos: {len(vencidos)}  ·  con riesgo: "
          f"{sum(r['Riesgo a la seguridad'] == 'Sí' for r in abiertos)}")
    print("Casos de prueba:")
    for r in reportes:
        if r["_caso"]:
            print(f"  {r['Folio']}  {r['_caso']:<38} {r['Tipo de falla']:<21} {r['Localidad']}")


def main():
    rng = random.Random(SEMILLA)
    municipios = cargar_municipios()
    poblacion, localidades = cargar_referencias(municipios)
    reportes = generar_reportes(rng, municipios, poblacion, localidades)

    (RAIZ / "datos").mkdir(exist_ok=True)
    (RAIZ / "sitio" / "descargas").mkdir(parents=True, exist_ok=True)
    escribir_reportes(reportes, RAIZ / "datos" / "reportes.xlsx")
    escribir_criterios(RAIZ / "datos" / "criterios.xlsx")
    for nombre in ("reportes.xlsx", "criterios.xlsx"):
        shutil.copyfile(RAIZ / "datos" / nombre, RAIZ / "sitio" / "descargas" / nombre)
    escribir_js(reportes, RAIZ / "sitio" / "datos" / "ejemplo.js")
    resumen(reportes)


if __name__ == "__main__":
    main()
